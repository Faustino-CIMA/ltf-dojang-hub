"""Club publication consent: one row per active member, and the exports."""

from __future__ import annotations

import csv
from datetime import date
from io import BytesIO, StringIO

from django.db import transaction
from django.db.models.functions import Lower
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response
from rest_framework.views import APIView

from clubs.models import Club
from members.models import Member

from licenses.pdf_utils import HTML, _bundled_logo_uri, _club_issuer, _invoice_sans_uris, _short_date

from .models import MemberRecord
from .views import ClubMgmtPermission, _require_club

CHANNELS = (
    ("publish_facebook", "Facebook"),
    ("publish_instagram", "Instagram"),
    ("publish_x", "X"),
    ("publish_tiktok", "TikTok"),
    ("publish_webpage", "Website"),
    ("publish_print", "Print media"),
)
CHANNEL_FIELDS = {field for field, _label in CHANNELS}
EXPORT_HEADERS = ["Last name", "First name", "Date of birth", "Age", *[label for _field, label in CHANNELS]]
PDF_RULES = ("all", "allowed", "denied", "denied_any")


def _age(born: date | None, on: date) -> int | None:
    if born is None:
        return None
    return on.year - born.year - ((on.month, on.day) < (born.month, born.day))


def _flags(record: MemberRecord | None) -> dict:
    return {field: bool(getattr(record, field, False)) if record is not None else False for field, _label in CHANNELS}


def _row(member: Member, today: date) -> dict:
    record = getattr(member, "club_record", None)
    born = member.date_of_birth
    return {
        "member": member.id,
        "last_name": member.last_name,
        "first_name": member.first_name,
        "date_of_birth": born.isoformat() if born else None,
        "age": _age(born, today),
        **_flags(record),
    }


def publication_rows(club) -> list[dict]:
    today = timezone.localdate()
    members = (
        Member.objects.filter(club=club, is_active=True)
        .select_related("club_record")
        .order_by(Lower("last_name"), Lower("first_name"), "id")
    )
    return [_row(member, today) for member in members]


def update_publication(member: Member, payload) -> dict:
    updates = {}
    for field in CHANNEL_FIELDS:
        if field not in payload:
            continue
        value = payload.get(field)
        if not isinstance(value, bool):
            raise ValueError("Choose yes or no for each publication consent.")
        updates[field] = value
    if not updates:
        raise ValueError("Choose a publication consent.")
    record, _created = MemberRecord.objects.get_or_create(member=member)
    for field, value in updates.items():
        setattr(record, field, value)
    record.save(update_fields=list(updates))
    member.club_record = record
    return _row(member, timezone.localdate())


def update_publication_column(club, field, value) -> list[dict]:
    if field not in CHANNEL_FIELDS:
        raise ValueError("Choose a publication channel.")
    if not isinstance(value, bool):
        raise ValueError("Choose yes or no for each publication consent.")
    member_ids = list(Member.objects.filter(club=club, is_active=True).values_list("id", flat=True))
    now = timezone.now()
    with transaction.atomic():
        existing = list(MemberRecord.objects.filter(member_id__in=member_ids))
        existing_ids = {record.member_id for record in existing}
        missing = [
            MemberRecord(member_id=member_id, updated_at=now, **{field: value})
            for member_id in member_ids
            if member_id not in existing_ids
        ]
        if missing:
            MemberRecord.objects.bulk_create(missing)
        if existing:
            for record in existing:
                setattr(record, field, value)
                record.updated_at = now
            MemberRecord.objects.bulk_update(existing, [field, "updated_at"])
    return publication_rows(club)


def publication_channels(raw) -> list[tuple[str, str]]:
    if raw is None or str(raw).strip() == "":
        return []
    chosen: set[str] = set()
    for part in str(raw).split(","):
        name = part.strip()
        if not name:
            continue
        if name not in CHANNEL_FIELDS:
            raise ValueError("Choose a publication channel.")
        chosen.add(name)
    return [item for item in CHANNELS if item[0] in chosen]


def _channel_list(labels: list[str]) -> str:
    if len(labels) == 1:
        return labels[0]
    if len(labels) == 2:
        return f"{labels[0]} and {labels[1]}"
    return f"{', '.join(labels[:-1])}, and {labels[-1]}"


def filter_publication_rows(rows: list[dict], rule: str, channels: list[tuple[str, str]]) -> tuple[list[dict], str]:
    if rule not in PDF_RULES:
        raise ValueError("Choose which members to include.")
    if rule == "all":
        return rows, "All active members"
    if not channels:
        raise ValueError("Choose at least one publication channel.")
    fields = [field for field, _label in channels]
    labels = _channel_list([label for _field, label in channels])
    if rule == "allowed":
        kept = [row for row in rows if all(row[field] for field in fields)]
        scope = f"Members allowed on {labels}"
    elif rule == "denied":
        kept = [row for row in rows if all(not row[field] for field in fields)]
        scope = f"Members not allowed on {labels}"
    else:
        kept = [row for row in rows if any(not row[field] for field in fields)]
        scope = (
            f"Members not allowed on {labels}"
            if len(fields) == 1
            else f"Members not allowed on at least one of {labels}"
        )
    return kept, scope


def _yes_no(value: bool) -> str:
    return "Yes" if value else "No"


def _export_line(row: dict) -> list:
    return [
        row["last_name"],
        row["first_name"],
        row["date_of_birth"] or "",
        "" if row["age"] is None else row["age"],
        *[_yes_no(row[field]) for field, _label in CHANNELS],
    ]


def publication_csv(club) -> bytes:
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    for row in publication_rows(club):
        writer.writerow(_export_line(row))
    return buffer.getvalue().encode("utf-8-sig")


def publication_xlsx(club) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Publication consent"
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    header_fill = PatternFill("solid", fgColor="F2F4F7")
    header_font = Font(bold=True)
    sheet["A1"] = club.name
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"Publication consent · {_short_date(timezone.localdate())}"
    for column, value in enumerate(EXPORT_HEADERS, start=1):
        cell = sheet.cell(4, column, value)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for offset, row in enumerate(publication_rows(club)):
        values = _export_line(row)
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(5 + offset, column, value)
            cell.border = thin
            if column >= 3:
                cell.alignment = Alignment(horizontal="center")
    widths = [22, 18, 16, 8, 14, 14, 8, 12, 14, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"
    last_row = 4 + sum(1 for _row in sheet.iter_rows(min_row=5, max_col=1, values_only=True) if _row[0] not in (None, ""))
    sheet.auto_filter.ref = f"A4:J{max(4, last_row)}"
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def publication_pdf(club, *, base_url: str, rows: list[dict] | None = None, scope: str = "All active members") -> bytes | None:
    if HTML is None:
        return None
    today = timezone.localdate()
    prepared = []
    for row in publication_rows(club) if rows is None else rows:
        born = date.fromisoformat(row["date_of_birth"]) if row["date_of_birth"] else None
        prepared.append(
            {
                **row,
                "born_label": _short_date(born),
                "age_label": "" if row["age"] is None else str(row["age"]),
                "marks": [_yes_no(row[field]) for field, _label in CHANNELS],
            }
        )
    html = render_to_string(
        "finance/publication_consent_pdf.html",
        {
            "issuer": _club_issuer(club),
            "invoice_fonts": _invoice_sans_uris(),
            "footer_logos": {
                "world": _bundled_logo_uri("world-taekwondo.png"),
                "europe": _bundled_logo_uri("european-taekwondo.png"),
                "ltf": _bundled_logo_uri("ltf-letterhead-logo.jpg"),
            },
            "channels": [label for _field, label in CHANNELS],
            "rows": prepared,
            "printed_on": _short_date(today),
            "count": len(prepared),
            "scope": scope,
        },
    )
    return HTML(string=html, base_url=base_url).write_pdf()


def _club(request) -> Club:
    club_id = _require_club(request)
    club = Club.objects.filter(pk=club_id).first()
    if club is None:
        raise PermissionDenied(detail="This module is not available.")
    return club


def _attachment(payload: bytes, content_type: str, filename: str) -> HttpResponse:
    response = HttpResponse(payload, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


class PublicationConsentView(APIView):
    permission_classes = [ClubMgmtPermission]

    def get(self, request):
        return Response({"members": publication_rows(_club(request))})

    def post(self, request):
        club = _club(request)
        try:
            member_id = int(request.data.get("member"))
        except (TypeError, ValueError):
            return Response({"detail": "Choose a member."}, status=400)
        member = Member.objects.filter(pk=member_id, club=club, is_active=True).select_related("club_record").first()
        if member is None:
            return Response({"detail": "Member not found in this club."}, status=400)
        try:
            row = update_publication(member, request.data)
        except ValueError as error:
            return Response({"detail": str(error)}, status=400)
        return Response(row)


class PublicationConsentColumnView(APIView):
    permission_classes = [ClubMgmtPermission]

    def post(self, request):
        try:
            rows = update_publication_column(_club(request), request.data.get("field"), request.data.get("value"))
        except ValueError as error:
            return Response({"detail": str(error)}, status=400)
        return Response({"members": rows})


class PublicationConsentCsvView(APIView):
    permission_classes = [ClubMgmtPermission]

    def get(self, request):
        club = _club(request)
        return _attachment(publication_csv(club), "text/csv; charset=utf-8", f"publication-consent-{club.id}.csv")


class PublicationConsentXlsxView(APIView):
    permission_classes = [ClubMgmtPermission]

    def get(self, request):
        club = _club(request)
        return _attachment(
            publication_xlsx(club),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"publication-consent-{club.id}.xlsx",
        )


class PublicationConsentPdfView(APIView):
    permission_classes = [ClubMgmtPermission]

    def get(self, request):
        club = _club(request)
        rule = (request.query_params.get("rule") or "all").strip() or "all"
        try:
            channels = publication_channels(request.query_params.get("channels"))
            rows, scope = filter_publication_rows(publication_rows(club), rule, channels)
        except ValueError as error:
            return Response({"detail": str(error)}, status=400)
        pdf = publication_pdf(club, base_url=request.build_absolute_uri("/"), rows=rows, scope=scope)
        if not pdf:
            return Response({"detail": "PDF generation is not available."}, status=500)
        return _attachment(pdf, "application/pdf", f"publication-consent-{club.id}.pdf")
