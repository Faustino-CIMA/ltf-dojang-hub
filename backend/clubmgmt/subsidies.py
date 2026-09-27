from __future__ import annotations

import csv
import re
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

from django.db.models import Q
from django.utils import timezone

from clubs.models import FederationProfile
from licenses.models import License
from members.grades import OFFICIAL_GRADES
from members.models import Member

from .models import CoachQualification, Committee, CommitteeMandate, ExtraordinarySubsidy, SubsidySeason
from .training_models import TrainingSeries, TrainingSettings

FORMS = Path(__file__).resolve().parent / "subsidy_forms"
FEDERATION_FALLBACK = "Luxembourg Taekwondo Federation"
COACH_POINTS = {
    CoachQualification.Level.EQF1: 20,
    CoachQualification.Level.EQF2: 20,
    CoachQualification.Level.EQF2BIS: 40,
    CoachQualification.Level.EQF3: 40,
    CoachQualification.Level.EQF4: 60,
    CoachQualification.Level.EQF5: 100,
    CoachQualification.Level.EQF6: 100,
}
YOUTH_QUALIFIED = {
    CoachQualification.Level.EQF2BIS,
    CoachQualification.Level.EQF3,
    CoachQualification.Level.EQF4,
    CoachQualification.Level.EQF5,
    CoachQualification.Level.EQF6,
}
FORM_FILES = {
    ExtraordinarySubsidy.Kind.CHAMPIONSHIP: "demande-participation.pdf",
    ExtraordinarySubsidy.Kind.CUP: "demande-coupe.pdf",
    ExtraordinarySubsidy.Kind.ORGANISATION: "demande-organisation.pdf",
}


def included_in_qualite(qualification: CoachQualification | None) -> bool:
    return qualification is None or qualification.include_in_qualite


def age_on(born: date | None, on: date) -> int | None:
    if born is None:
        return None
    return on.year - born.year - ((on.month, on.day) < (born.month, born.day))


def is_adult(born: date | None, on: date) -> bool:
    age = age_on(born, on)
    return age is not None and age >= 18


def _adult_ids(members, cutoff: date) -> set[int]:
    return {member.id for member in members if is_adult(member.date_of_birth, cutoff)}


def _member_rows(members, cutoff: date) -> list[dict]:
    rows = []
    for member in members:
        age = age_on(member.date_of_birth, cutoff)
        rows.append(
            {
                "id": member.id,
                "name": f"{member.first_name} {member.last_name}".strip(),
                "under_16": age is not None and age < 16,
                "adult": age is not None and age >= 18,
            }
        )
    return rows


def default_deadline(year: int) -> date:
    return date(year, 9, 30)


def inaps_deadline(year: int) -> date:
    return date(year, 12, 15)


def diploma_deadline(year: int) -> date:
    return date(year + 1, 2, 1)


def plus_three_months(value: date) -> date:
    month = value.month + 3
    year = value.year + (month - 1) // 12
    month = (month - 1) % 12 + 1
    day = min(value.day, 28)
    return date(year, month, day)


def get_or_create_season(club, year: int) -> SubsidySeason:
    season, _created = SubsidySeason.objects.get_or_create(
        club=club,
        year=year,
        defaults={"deadline": default_deadline(year)},
    )
    return season


def _licensed_ids(club_id: int, year: int) -> set[int]:
    return set(
        License.objects.filter(
            club_id=club_id,
            year=year,
            status__in=[License.Status.ACTIVE, License.Status.PENDING],
        ).values_list("member_id", flat=True)
    )


def _officer_name(club, roles: list[str], *, formal: bool = False) -> str:
    today = timezone.localdate()
    mandate = (
        CommitteeMandate.objects.filter(
            committee__scope=Committee.Scope.CLUB,
            committee__club=club,
            role__in=roles,
            started_on__lte=today,
        )
        .filter(Q(ended_on__isnull=True) | Q(ended_on__gte=today))
        .select_related("member", "person")
        .first()
    )
    if mandate is None:
        return ""
    if mandate.member_id:
        first, last = mandate.member.first_name, mandate.member.last_name
        return _formal_name(last, first) if formal else f"{first} {last}".strip()
    if mandate.person_id:
        first, last = mandate.person.first_name, mandate.person.last_name
        return _formal_name(last, first) if formal else f"{first} {last}".strip()
    return mandate.title


def volunteer_points(headcount: int) -> int:
    if headcount < 50:
        return 150
    if headcount <= 200:
        return 300
    return 500


def subsidy_income_category(club):
    from licenses.models import IncomeCategory

    category, _created = IncomeCategory.objects.get_or_create(
        club=club,
        code="subsidy",
        defaults={"name": "Subsidies", "sort_order": 40, "is_active": True},
    )
    return category


def sync_subsidy_income(season: SubsidySeason, actor, *, update_existing: bool = False) -> str:
    from licenses.models import Income, Order

    from .access import can_record_club_payments

    amount = season.paid_amount
    paid = season.status == SubsidySeason.Status.PAID and amount is not None and Decimal(amount) > 0
    if season.income_id and season.income.status == Income.Status.VOID:
        season.income = None
        season.save(update_fields=["income"])
    if not paid:
        return "recorded" if season.income_id else "skipped"
    if season.income_id and not update_existing:
        return "recorded"
    if not can_record_club_payments(actor, season.club_id):
        return "recorded" if season.income_id else "needs_officer"
    paid_on = season.paid_on or timezone.localdate()
    if season.income_id:
        income = season.income
        income.amount = amount
        income.income_date = paid_on
        income.description = f"Sports subsidy {season.year}"
        income.payer = "Ministère des Sports"
        income.ledger = Order.Ledger.CLUB
        income.save()
        return "updated"
    income = Income(
        category=subsidy_income_category(season.club),
        club=season.club,
        description=f"Sports subsidy {season.year}",
        payer="Ministère des Sports",
        amount=amount,
        income_date=paid_on,
        received_at=timezone.now(),
        status=Income.Status.RECEIVED,
        payment_method="bank_transfer",
        reference=f"subsidy-{season.club_id}-{season.year}",
        ledger=Order.Ledger.CLUB,
        created_by=actor if actor and getattr(actor, "is_authenticated", False) else None,
    )
    income.save()
    season.income = income
    season.save(update_fields=["income"])
    return "recorded"


def build_dossier(club, year: int) -> dict:
    season = get_or_create_season(club, year)
    cutoff = date(year, 12, 31)
    members = list(Member.objects.filter(club=club, is_active=True).select_related("club_record"))
    licensed = _licensed_ids(club.id, year)
    bands = {
        "male_under_16": 0,
        "female_under_16": 0,
        "male_16_plus": 0,
        "female_16_plus": 0,
    }
    youth = []
    for member in members:
        if member.id not in licensed:
            continue
        years = age_on(member.date_of_birth, cutoff)
        under = years is not None and years < 16
        female = member.sex == Member.Sex.FEMALE
        if under:
            bands["female_under_16" if female else "male_under_16"] += 1
        else:
            bands["female_16_plus" if female else "male_16_plus"] += 1
        if under and _is_competition_athlete(member):
            record = getattr(member, "club_record", None)
            youth.append(
                {
                    "id": member.id,
                    "first_name": member.first_name,
                    "last_name": member.last_name,
                    "date_of_birth": member.date_of_birth.isoformat() if member.date_of_birth else "",
                    "sex": member.sex,
                    "licence": member.ltf_licenseid or member.wt_licenseid,
                    "national_id": (record.social_security_number if record else "") or "",
                }
            )
    licensed_total = sum(bands.values())
    headcount = licensed_total + season.non_licensed_count
    qualifications = {
        row.user_id: row
        for row in CoachQualification.objects.filter(club=club)
    }
    coaches = []
    coach_points = 0
    youth_coaches = 0
    youth_qualified = 0
    for user in club.trainers.select_related("member_profile").order_by("last_name", "first_name"):
        qualification = qualifications.get(user.id)
        if not included_in_qualite(qualification):
            continue
        level = qualification.eqf_level if qualification else ""
        under = bool(qualification and qualification.coaches_under_16)
        diploma = qualification.diploma_status if qualification else CoachQualification.Diploma.MISSING
        has_diploma = bool(qualification and qualification.diploma_file)
        points = COACH_POINTS.get(level, 0)
        coach_points += points
        if under:
            youth_coaches += 1
            if level in YOUTH_QUALIFIED:
                youth_qualified += 1
        member = getattr(user, "member_profile", None)
        coaches.append(
            {
                "user_id": user.id,
                "name": f"{(member.first_name if member else user.first_name) or ''} {(member.last_name if member else user.last_name) or ''}".strip()
                or user.username,
                "eqf_level": level,
                "coaches_under_16": under,
                "diploma_status": diploma,
                "has_diploma": has_diploma,
                "diploma_name": Path(qualification.diploma_file.name).name if has_diploma else "",
                "points": points,
            }
        )
    half_ok = youth_coaches == 0 or (youth_qualified * 2 >= youth_coaches)
    has_eqf3 = any(row["eqf_level"] in YOUTH_QUALIFIED for row in coaches)
    diplomas_ok = bool(coaches) and all(
        row["diploma_status"] != CoachQualification.Diploma.MISSING or row["has_diploma"] for row in coaches
    )
    checklist = [
        {"id": "affiliated", "ok": bool(club.is_active)},
        {"id": "season", "ok": season.season_complete},
        {"id": "youth", "ok": bands["male_under_16"] + bands["female_under_16"] > 0},
        {"id": "eqf3", "ok": has_eqf3},
        {"id": "half", "ok": half_ok and youth_coaches > 0},
        {"id": "diplomas", "ok": diplomas_ok},
        {"id": "iban", "ok": bool(club.iban)},
        {"id": "rib", "ok": season.rib_attached or bool(season.rib_file)},
        {"id": "myguichet", "ok": season.myguichet_users >= 2},
    ]
    return {
        "year": year,
        "deadline": season.deadline.isoformat(),
        "dates": {
            "file_by": season.deadline.isoformat(),
            "inaps_by": inaps_deadline(year).isoformat(),
            "diplomas_by": diploma_deadline(year).isoformat(),
        },
        "season": {
            "season_complete": season.season_complete,
            "myguichet_users": season.myguichet_users,
            "rib_attached": season.rib_attached,
            "has_rib": bool(season.rib_file),
            "rib_name": Path(season.rib_file.name).name if season.rib_file else "",
            "non_licensed_count": season.non_licensed_count,
            "status": season.status,
            "submitted_on": season.submitted_on.isoformat() if season.submitted_on else "",
            "paid_on": season.paid_on.isoformat() if season.paid_on else "",
            "paid_amount": str(season.paid_amount) if season.paid_amount is not None else "",
            "income_number": season.income.income_number if season.income_id else "",
            "income_needs_officer": bool(
                season.status == SubsidySeason.Status.PAID
                and season.paid_amount
                and Decimal(season.paid_amount) > 0
                and not season.income_id
            ),
        },
        "checklist": checklist,
        "effectifs": {
            **bands,
            "licensed_total": licensed_total,
            "non_licensed": season.non_licensed_count,
            "headcount": headcount,
            "volunteer_points": volunteer_points(headcount),
        },
        "coach_points": coach_points,
        "eligible_youth": len(youth),
        "qualite_estimate": f"{Decimal(len(youth) * 150):.2f}",
        "training_place": TrainingSettings.objects.filter(club=club).values_list("default_place", flat=True).first() or "",
        "coaches": coaches,
        "headcount_rows": headcount_rows(club, year, season, members, licensed),
        "youth": youth,
        "members": _member_rows(members, cutoff),
        "cases": [
            {
                "id": row.id,
                "kind": row.kind,
                "title": row.title,
                "place": row.place,
                "starts_on": row.starts_on.isoformat() if row.starts_on else "",
                "ends_on": row.ends_on.isoformat() if row.ends_on else "",
                "status": row.status,
                "athlete_ids": row.athlete_ids,
                "official_ids": [
                    member_id
                    for member_id in row.official_ids or []
                    if member_id in _adult_ids(members, cutoff)
                ],
                "travel_mode": row.travel_mode,
                "travel_units": str(row.travel_units) if row.travel_units is not None else "",
                "travel_rate": str(row.travel_rate) if row.travel_rate is not None else "",
                "stay_people": row.stay_people,
                "stay_days": row.stay_days,
                "stay_rate": str(row.stay_rate) if row.stay_rate is not None else "",
                "entry_fee": str(row.entry_fee) if row.entry_fee is not None else "",
                "medical_fee": str(row.medical_fee) if row.medical_fee is not None else "",
                "supplies_fee": str(row.supplies_fee) if row.supplies_fee is not None else "",
                "notes": row.notes,
                "account_due": plus_three_months(row.ends_on or row.starts_on).isoformat()
                if (row.ends_on or row.starts_on)
                else "",
            }
            for row in club.extraordinary_subsidies.filter(year=year)
        ],
    }


def _sex_key(member) -> str:
    return "F" if member.sex == Member.Sex.FEMALE else "M"


def _age_band(years: int | None) -> str:
    if years is None:
        return "unknown"
    if years < 16:
        return "under_16"
    if years <= 17:
        return "youth_16"
    if years <= 34:
        return "senior"
    return "master"


def _roles(member) -> set[str]:
    return {member.primary_license_role, member.secondary_license_role} - {""}


def _is_competition_athlete(member) -> bool:
    roles = _roles(member)
    return not roles or Member.LicenseRole.ATHLETE in roles


def headcount_rows(club, year: int, season: SubsidySeason, members, licensed: set[int]) -> list[dict]:
    """Rows shaped like the MyGuichet effectifs page.

    Ages are on 31 December. Youth over 16 are 16–17, seniors 18–34, and
    veterans/masters 35 and over. Competition rows count a licensed member
    whose role is Athlete or unset. Other rows count that role even when the
    person is also an athlete, and they are not added again into the club total.
    """
    cutoff = date(year, 12, 31)
    age_counts = {key: {"M": 0, "F": 0} for key in ("under_16", "youth_16", "senior", "master", "unknown")}
    other = {key: {"M": 0, "F": 0} for key in ("leisure", "officials", "referees", "coaches", "other")}
    for member in members:
        if member.id not in licensed:
            continue
        sex = _sex_key(member)
        roles = _roles(member)
        if _is_competition_athlete(member):
            age_counts[_age_band(age_on(member.date_of_birth, cutoff))][sex] += 1
        if Member.LicenseRole.FAN in roles:
            other["leisure"][sex] += 1
        if Member.LicenseRole.OFFICIAL in roles:
            other["officials"][sex] += 1
        if Member.LicenseRole.REFEREE in roles:
            other["referees"][sex] += 1
        if Member.LicenseRole.COACH in roles:
            other["coaches"][sex] += 1
        extra = roles - {
            Member.LicenseRole.ATHLETE,
            Member.LicenseRole.FAN,
            Member.LicenseRole.OFFICIAL,
            Member.LicenseRole.REFEREE,
            Member.LicenseRole.COACH,
        }
        if extra:
            other["other"][sex] += 1

    def count(pair, label):
        male, female = pair["M"], pair["F"]
        return {"kind": "count", "label": label, "male": male, "female": female, "total": male + female}

    def total(pairs, label):
        male = sum(pair["M"] for pair in pairs)
        female = sum(pair["F"] for pair in pairs)
        return {"kind": "total", "label": label, "male": male, "female": female, "total": male + female}

    age_pairs = [age_counts[key] for key in ("under_16", "youth_16", "senior", "master", "unknown")]
    other_pairs = [other[key] for key in ("leisure", "officials", "referees", "coaches", "other")]
    licensed_sub = total(age_pairs, "subsidiesLicensedSubtotal")
    non_licensed = season.non_licensed_count
    return [
        {"kind": "section", "label": "subsidiesLicensedCategories", "male": None, "female": None, "total": None},
        count(age_counts["under_16"], "subsidiesAgeUnder16"),
        count(age_counts["youth_16"], "subsidiesAgeYouth16"),
        count(age_counts["senior"], "subsidiesAgeSenior"),
        count(age_counts["master"], "subsidiesAgeMaster"),
        count(age_counts["unknown"], "subsidiesAgeUnknown"),
        licensed_sub,
        {"kind": "section", "label": "subsidiesOtherLicences", "male": None, "female": None, "total": None},
        count(other["leisure"], "subsidiesLeisure"),
        count(other["officials"], "subsidiesOfficials"),
        count(other["referees"], "subsidiesReferees"),
        count(other["coaches"], "subsidiesCoachLicences"),
        count(other["other"], "subsidiesOtherRoles"),
        total(other_pairs, "subsidiesOtherSubtotal"),
        {"kind": "section", "label": "subsidiesNonLicensedSection", "male": None, "female": None, "total": None},
        {
            "kind": "count",
            "label": "subsidiesNonLicensedRow",
            "male": None,
            "female": None,
            "total": non_licensed,
        },
        {
            "kind": "total",
            "label": "subsidiesClubTotal",
            "male": licensed_sub["male"],
            "female": licensed_sub["female"],
            "total": licensed_sub["total"] + non_licensed,
        },
    ]


_FRENCH_MONTHS = (
    "janvier",
    "février",
    "mars",
    "avril",
    "mai",
    "juin",
    "juillet",
    "août",
    "septembre",
    "octobre",
    "novembre",
    "décembre",
)
_EQF_FR = {
    CoachQualification.Level.EQF1: "EQF 1",
    CoachQualification.Level.EQF2: "EQF 2",
    CoachQualification.Level.EQF2BIS: "EQF 2 bis",
    CoachQualification.Level.EQF3: "EQF 3",
    CoachQualification.Level.EQF4: "EQF 4",
    CoachQualification.Level.EQF5: "EQF 5",
    CoachQualification.Level.EQF6: "EQF 6",
}


def _french_date(value: date) -> str:
    return f"{value.day} {_FRENCH_MONTHS[value.month - 1]} {value.year}"


def _formal_name(last: str, first: str) -> str:
    last_name = (last or "").strip()
    first_name = (first or "").strip()
    if last_name:
        return f"{last_name.upper()} {first_name}".strip()
    return first_name


def _club_logo_uri(club) -> str:
    from clubs.models import BrandingAsset

    logos = list(
        BrandingAsset.objects.filter(
            club=club,
            asset_type=BrandingAsset.AssetType.LOGO,
            is_selected=True,
        )
    )
    preferred = (
        BrandingAsset.UsageType.PRINT,
        BrandingAsset.UsageType.GENERAL,
        BrandingAsset.UsageType.INVOICE,
        BrandingAsset.UsageType.DIGITAL,
    )
    chosen = None
    for usage in preferred:
        chosen = next((row for row in logos if row.usage_type == usage), None)
        if chosen:
            break
    if chosen is None and logos:
        chosen = logos[0]
    if chosen is None or not chosen.file:
        return ""
    try:
        return Path(chosen.file.path).resolve().as_uri()
    except (NotImplementedError, ValueError, OSError):
        return ""


def _address_lines(club) -> list[str]:
    lines = [line for line in (club.address_line1, club.address_line2) if str(line or "").strip()]
    if not lines and club.address:
        lines.append(club.address)
    postal = " ".join(part for part in (club.postal_code, club.locality or club.city) if str(part or "").strip())
    if postal:
        lines.append(postal)
    elif club.city and club.city not in lines:
        lines.append(club.city)
    return lines


_GRADE_RANK = {grade.casefold(): index for index, grade in enumerate(OFFICIAL_GRADES)}
_GRADE_WORDS = re.compile(r"(\d+)\s*(?:ere|er|e|st|nd|rd|th)?\s*(kup|poom|dan)", re.IGNORECASE)


def _grade_rank(grade: str) -> int:
    """Higher official grade, higher number. Unknown grades sort last."""
    text = (grade or "").strip().casefold().replace("°", " ").replace("è", "e").replace("é", "e")
    if text in _GRADE_RANK:
        return _GRADE_RANK[text]
    match = _GRADE_WORDS.search(text)
    if match is None:
        return -1
    number = int(match.group(1))
    family = match.group(2).casefold()
    if number == 1:
        ordinal = "1st"
    elif number == 2:
        ordinal = "2nd"
    elif number == 3:
        ordinal = "3rd"
    else:
        ordinal = f"{number}th"
    return _GRADE_RANK.get(f"{ordinal} {family}".casefold(), -1)


def _by_grade_desc(grade: str, name: str) -> tuple:
    return (-(_grade_rank(grade) + 1), name.casefold())


def render_trainers_list_pdf(club, year: int) -> bytes:
    from django.template.loader import render_to_string

    from licenses.pdf_utils import HTML

    qualifications = {row.user_id: row for row in CoachQualification.objects.filter(club=club)}
    trainers = []
    for user in club.trainers.select_related("member_profile").order_by("last_name", "first_name"):
        qualification = qualifications.get(user.id)
        if not included_in_qualite(qualification):
            continue
        member = getattr(user, "member_profile", None)
        last = (member.last_name if member else "") or user.last_name
        first = (member.first_name if member else "") or user.first_name
        under = ""
        level = ""
        if qualification:
            level = _EQF_FR.get(qualification.eqf_level, "")
            under = "Oui" if qualification.coaches_under_16 else "Non"
        trainers.append(
            {
                "number": len(trainers) + 1,
                "name": _formal_name(last, first) or user.username,
                "grade": (member.belt_rank if member else "") or "",
                "qualification": level,
                "under_16": under,
            }
        )
    trainers.sort(key=lambda row: _by_grade_desc(row["grade"], row["name"]))
    for index, row in enumerate(trainers, start=1):
        row["number"] = index
    place = (club.locality or club.city or "").strip()
    today = _french_date(timezone.localdate())
    html = render_to_string(
        "subsidies/trainers_list.html",
        {
            "club_name": club.name,
            "address_lines": _address_lines(club),
            "email": club.email,
            "logo_uri": _club_logo_uri(club),
            "year": year,
            "place_date": f"{place}, le {today}" if place else f"Le {today}",
            "trainers": trainers,
            "president": _officer_name(club, [CommitteeMandate.Role.PRESIDENT], formal=True),
        },
    )
    if HTML is None:
        raise RuntimeError("PDF renderer is not available.")
    return HTML(string=html, base_url=str(Path(__file__).resolve().parents[1])).write_pdf()


_WEEKDAYS_FR = ("Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche")


def _qualite_coach_names(series: TrainingSeries, qualifications: dict) -> str:
    names = []
    for user in series.coaches.all():
        if not included_in_qualite(qualifications.get(user.id)):
            continue
        member = getattr(user, "member_profile", None)
        last = (member.last_name if member else "") or user.last_name
        first = (member.first_name if member else "") or user.first_name
        grade = (member.belt_rank if member else "") or ""
        names.append((_formal_name(last, first) or user.username, grade))
    names.sort(key=lambda item: _by_grade_desc(item[1], item[0]))
    return ", ".join(name for name, _grade in names)


def training_list_grid(club, year: int) -> tuple[list[str], list[dict], int]:
    """Weekly overview. Columns are only the weekdays that have a class."""
    start = date(year, 1, 1)
    end = date(year, 12, 31)
    qualifications = {row.user_id: row for row in CoachQualification.objects.filter(club=club)}
    classes = (
        TrainingSeries.objects.filter(club=club, active=True, valid_from__lte=end, valid_until__gte=start)
        .prefetch_related("coaches__member_profile")
        .order_by("start_time", "end_time", "weekday", "name", "id")
    )
    slots: dict[tuple, list[list[dict]]] = {}
    used: set[int] = set()
    class_count = 0
    for series in classes:
        if not 0 <= series.weekday <= 6:
            continue
        key = (series.start_time, series.end_time)
        days = slots.setdefault(key, [[] for _ in range(7)])
        days[series.weekday].append(
            {
                "category": series.name,
                "coaches": _qualite_coach_names(series, qualifications),
            }
        )
        used.add(series.weekday)
        class_count += 1
    weekday_indexes = [index for index in range(7) if index in used]
    rows = []
    for (start_time, end_time), days in slots.items():
        rows.append(
            {
                "time": f"{start_time.strftime('%H:%M')} - {end_time.strftime('%H:%M')}",
                "days": [days[index] for index in weekday_indexes],
            }
        )
    return [_WEEKDAYS_FR[index] for index in weekday_indexes], rows, class_count


def render_training_list_pdf(club, year: int) -> bytes:
    from django.template.loader import render_to_string

    from licenses.pdf_utils import HTML

    weekdays, rows, class_count = training_list_grid(club, year)
    place = (club.locality or club.city or "").strip()
    today = _french_date(timezone.localdate())
    html = render_to_string(
        "subsidies/training_list.html",
        {
            "club_name": club.name,
            "address_lines": _address_lines(club),
            "email": club.email,
            "logo_uri": _club_logo_uri(club),
            "year": year,
            "place_date": f"{place}, le {today}" if place else f"Le {today}",
            "weekdays": weekdays,
            "rows": rows,
            "class_count": class_count,
            "president": _officer_name(club, [CommitteeMandate.Role.PRESIDENT], formal=True),
        },
    )
    if HTML is None:
        raise RuntimeError("PDF renderer is not available.")
    return HTML(string=html, base_url=str(Path(__file__).resolve().parents[1])).write_pdf()


def youth_csv(club, year: int) -> str:
    dossier = build_dossier(club, year)
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        ["First name", "Last name", "Date of birth", "Sex", "Licence", "National identification number"]
    )
    for row in dossier["youth"]:
        writer.writerow(
            [
                row["first_name"],
                row["last_name"],
                row["date_of_birth"],
                row["sex"],
                row["licence"],
                row["national_id"],
            ]
        )
    return buffer.getvalue()


_HEADCOUNT_LABELS = {
    "subsidiesLicensedCategories": "1) Licences par catégories",
    "subsidiesAgeUnder16": "Jeunes < 16 ans",
    "subsidiesAgeYouth16": "Jeunes > 16 ans (16–17)",
    "subsidiesAgeSenior": "Seniors (18–34)",
    "subsidiesAgeMaster": "Vétérans / Masters (35+)",
    "subsidiesAgeUnknown": "Licenciés sans date de naissance",
    "subsidiesLicensedSubtotal": "SOUS-TOTAL",
    "subsidiesOtherLicences": "2) Autres licences",
    "subsidiesLeisure": "Licences loisirs",
    "subsidiesOfficials": "Dirigeants administratifs",
    "subsidiesReferees": "Arbitres / Juges",
    "subsidiesCoachLicences": "Entraîneurs / Moniteurs",
    "subsidiesOtherRoles": "Autres rôles (hors formulaire)",
    "subsidiesOtherSubtotal": "SOUS-TOTAL",
    "subsidiesNonLicensedSection": "3) Non licenciés du club",
    "subsidiesNonLicensedRow": "Effectifs non licenciés du club",
    "subsidiesClubTotal": "TOTAL DES EFFECTIFS DU CLUB",
}


def youth_xlsx(club, year: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    dossier = build_dossier(club, year)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Effectifs"
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    header_fill = PatternFill("solid", fgColor="F2F4F7")
    section_font = Font(bold=True)
    sheet["A1"] = club.name
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"Effectifs {year} · âges au 31 décembre"
    sheet["A3"] = (
        "Jeunes > 16 ans = 16–17 ans. Seniors = 18–34 ans. Vétérans / Masters = 35 ans et plus. "
        "Les autres licences ne sont pas ajoutées au total : une personne déjà comptée comme athlète y figure une seconde fois. "
        "Les non-licenciés sont un seul total, sans répartition par sexe."
    )
    sheet["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A3:D3")
    sheet.row_dimensions[3].height = 48
    headers = ["", "Masculins", "Féminins", "TOTAL"]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(5, column, value)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(horizontal="center")
    for offset, row in enumerate(dossier["headcount_rows"]):
        line = 6 + offset
        label = _HEADCOUNT_LABELS.get(row["label"], row["label"])
        sheet.cell(line, 1, label).border = thin
        if row["kind"] == "section":
            sheet.cell(line, 1).font = section_font
            for column in range(2, 5):
                sheet.cell(line, column).border = thin
                sheet.cell(line, column).fill = header_fill
            continue
        sheet.cell(line, 1).font = section_font if row["kind"] == "total" else Font()
        for column, key in ((2, "male"), (3, "female")):
            value = row[key]
            if row["kind"] == "count" and value == 0:
                value = None
            cell = sheet.cell(line, column, value)
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")
        total_cell = sheet.cell(line, 4)
        total_cell.border = thin
        total_cell.alignment = Alignment(horizontal="center")
        total_cell.font = section_font if row["kind"] == "total" else Font()
        if row["male"] is None and row["female"] is None:
            total_cell.value = row["total"] or None
        else:
            total_cell.value = f"=B{line}+C{line}"
    widths = (42, 14, 14, 12)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    names = workbook.create_sheet("Jeunes moins de 16")
    name_headers = ["Nom", "Prénom", "Date de naissance", "Sexe", "Licence", "Matricule national"]
    for column, value in enumerate(name_headers, start=1):
        cell = names.cell(1, column, value)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = thin
    for offset, row in enumerate(dossier["youth"], start=2):
        values = [
            row["last_name"],
            row["first_name"],
            row["date_of_birth"],
            "F" if row["sex"] == Member.Sex.FEMALE else "M",
            row["licence"],
            row["national_id"],
        ]
        for column, value in enumerate(values, start=1):
            cell = names.cell(offset, column, value)
            cell.border = thin
    for index, width in enumerate((22, 18, 20, 10, 16, 24), start=1):
        names.column_dimensions[get_column_letter(index)].width = width
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _fmt_date(value: date | None) -> str:
    return value.strftime("%d.%m.%Y") if value else ""


def _amount(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    return Decimal(value)


def _fmt_amount(value: Decimal) -> str:
    if not value:
        return ""
    return f"{value:.2f}"


def _person_name(member) -> str:
    return f"{member.last_name} {member.first_name}".strip()


def _divers_lines(case: ExtraordinarySubsidy) -> list[tuple[str, Decimal]]:
    rows = []
    if _amount(case.entry_fee):
        rows.append(("Inscription", _amount(case.entry_fee)))
    if _amount(case.medical_fee):
        rows.append(("Frais medicaux", _amount(case.medical_fee)))
    if _amount(case.supplies_fee):
        rows.append(("Fournitures", _amount(case.supplies_fee)))
    return rows


def _draw_overlay(page, values: dict) -> None:
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    annots = page.get("/Annots") or []
    pairs = []
    for annot in annots:
        obj = annot.get_object()
        name = str(obj.get("/T") or "")
        text = values.get(name)
        if text and obj.get("/Rect"):
            pairs.append((obj["/Rect"], text))
    if not pairs:
        return
    width = float(page.mediabox.width)
    height = float(page.mediabox.height)
    buffer = BytesIO()
    drawing = canvas.Canvas(buffer, pagesize=(width, height))
    drawing.setFont("Helvetica", 8)
    for rect, text in pairs:
        x1, y1, x2, y2 = [float(item) for item in rect]
        left = min(x1, x2) + 2
        bottom = min(y1, y2)
        top = max(y1, y2)
        lines = [line for line in str(text).split("\n") if line]
        if top - bottom > 28 and len(lines) > 1:
            cursor = top - 10
            for line in lines[:8]:
                if cursor < bottom:
                    break
                drawing.drawString(left, cursor, line[:90])
                cursor -= 10
        else:
            drawing.drawString(left, bottom + 2, " ".join(lines)[:110])
    drawing.save()
    buffer.seek(0)
    page.merge_page(PdfReader(buffer).pages[0])


def fill_extraordinary_pdf(case: ExtraordinarySubsidy) -> bytes:
    from pypdf import PdfReader, PdfWriter

    club = case.club
    federation = FederationProfile.objects.order_by("id").values_list("name", flat=True).first() or FEDERATION_FALLBACK
    suffix = {
        ExtraordinarySubsidy.Kind.CHAMPIONSHIP: "_2",
        ExtraordinarySubsidy.Kind.CUP: "_3",
        ExtraordinarySubsidy.Kind.ORGANISATION: "_4",
    }[case.kind]
    athletes = list(Member.objects.filter(id__in=case.athlete_ids or [], club=club).order_by("last_name", "first_name"))
    officials = [
        row
        for row in Member.objects.filter(id__in=case.official_ids or [], club=club).order_by("last_name", "first_name")
        if is_adult(row.date_of_birth, date(case.year, 12, 31))
    ]
    president = _officer_name(club, [CommitteeMandate.Role.PRESIDENT])
    mode = case.travel_mode or (ExtraordinarySubsidy.Travel.TRAIN if case.travel_units is not None and case.travel_rate is not None else "")
    travel_total = _amount(case.travel_units) * _amount(case.travel_rate) if mode else Decimal("0")
    stay_total = _amount(case.stay_people) * _amount(case.stay_days) * _amount(case.stay_rate)
    divers = _divers_lines(case)
    grand = travel_total + stay_total + sum((amount for _label, amount in divers), Decimal("0"))
    athlete_names = [_person_name(row) for row in athletes]
    official_lines = "\n".join(_person_name(row) for row in officials)
    overflow = athlete_names[3:]
    notes = (case.notes or "").strip()
    if case.status == ExtraordinarySubsidy.Status.ACCOUNTED:
        notes = "\n".join(part for part in ("Decompte final.", notes) if part)
    if overflow:
        extra = "Athletes: " + ", ".join(overflow)
        notes = "\n".join(part for part in (notes, extra) if part)
    values = {
        f"Fédération{suffix}": federation,
        f"Président{suffix}": president,
        f"Secrétaire id{suffix}": _officer_name(
            club, [CommitteeMandate.Role.SECRETARY, CommitteeMandate.Role.SECRETARY_GENERAL]
        ),
        f"Trésorier id{suffix}": _officer_name(club, [CommitteeMandate.Role.TREASURER]),
        f"Cpte IBAN No{suffix}": club.iban,
        f"4_{8 if case.kind == ExtraordinarySubsidy.Kind.CHAMPIONSHIP else 11 if case.kind == ExtraordinarySubsidy.Kind.CUP else 16}": club.bank_name,
    }
    if case.kind == ExtraordinarySubsidy.Kind.CHAMPIONSHIP:
        values.update(_championship_values(case, athletes, officials, athlete_names, official_lines, notes, mode, travel_total, stay_total, divers, grand))
    elif case.kind == ExtraordinarySubsidy.Kind.CUP:
        values.update(_cup_values(case, club, athletes, officials, athlete_names, notes, mode, travel_total, stay_total, divers, grand))
    else:
        values.update(_organisation_values(case, athletes, officials, athlete_names, notes, travel_total, stay_total, divers, grand, president))
    reader = PdfReader(str(FORMS / FORM_FILES[case.kind]))
    writer = PdfWriter()
    writer.append(reader)
    for page in writer.pages:
        writer.update_page_form_field_values(page, {key: value for key, value in values.items() if value}, auto_regenerate=False)
        _draw_overlay(page, values)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _put_travel(values, mode, units, rate, total, fields):
    spec = fields.get(mode)
    if not spec or not total:
        return
    units_field, rate_field, total_field = spec
    values[units_field] = _fmt_amount(_amount(units))
    values[rate_field] = _fmt_amount(_amount(rate))
    values[total_field] = _fmt_amount(total)


def _championship_values(case, athletes, officials, athlete_names, official_lines, notes, mode, travel_total, stay_total, divers, grand):
    values = {
        "I RENSEIGNEMENTS GENERAUX": case.title,
        "Motivation de la participation": case.notes,
        "undefined_15": case.place,
        "undefined_16": _fmt_date(case.starts_on),
        "au_8": _fmt_date(case.ends_on),
        "Départ : le": _fmt_date(case.starts_on),
        "Retour : le_2": _fmt_date(case.ends_on),
        "undefined_17": "1" if athletes else "",
        "Equipes": str(len(athletes)) if athletes else "",
        "athlètes": str(len(officials)) if officials else "",
        "lndiquer les noms et les": official_lines,
        "Remarques": notes,
        "Total": _fmt_amount(grand),
        "le_11": _fmt_date(timezone.localdate()),
        "Le Président ou son représentant_6": _officer_name(case.club, [CommitteeMandate.Role.PRESIDENT]),
    }
    name_fields = ["1_23", "2_23", "3_18"]
    perf_fields = [
        "réalisée dans cette discipline pendant la saison écoulée 1",
        "réalisée dans cette discipline pendant la saison écoulée 2",
        "réalisée dans cette discipline pendant la saison écoulée 3",
    ]
    for index, name in enumerate(athlete_names[:3]):
        values[name_fields[index]] = name
        values[perf_fields[index]] = "Taekwondo"
    _put_travel(
        values,
        mode,
        case.travel_units,
        case.travel_rate,
        travel_total,
        {
            "train": ("par train", "x_6", "en € 1"),
            "car": ("par voiture", "km à", "en € 2"),
            "plane": ("par avion", "x_7", "en € 3"),
        },
    )
    if stay_total:
        values["Frais de séjour 1"] = str(case.stay_people or "")
        values["x_8"] = str(case.stay_days or "")
        values["jours à"] = _fmt_amount(_amount(case.stay_rate))
        values["en € 4"] = _fmt_amount(stay_total)
    desc_fields = ["Frais de séjour 2", "Frais de séjour 3"]
    amount_fields = ["en € 5", "en € 6"]
    for index, (label, amount) in enumerate(divers[:2]):
        values[desc_fields[index]] = label
        values[amount_fields[index]] = _fmt_amount(amount)
    if len(divers) > 2:
        extra = "; ".join(f"{label} {_fmt_amount(amount)}" for label, amount in divers[2:])
        values["Remarques"] = "\n".join(part for part in (values.get("Remarques"), extra) if part)
    return values


def _cup_values(case, club, athletes, officials, athlete_names, notes, mode, travel_total, stay_total, divers, grand):
    financing = notes
    if athlete_names:
        financing = "\n".join(part for part in (financing, "Athletes: " + ", ".join(athlete_names)) if part)
    values = {
        "Club": club.name,
        "Dénomination de la coupe et de lorganisme international organisateur": case.title,
        "1ère rencontre  de finales, Lieu": case.place,
        "1ère rencontre  de finales, Date": _fmt_date(case.starts_on),
        "1ère rencontre  de finales, Nbre de vos participants actifs": str(len(athletes)) if athletes else "",
        "1ère rencontre  de finales, Nbre de vos officiels": str(len(officials)) if officials else "",
        "Indiquez le mode de financement des rencontres et le cas échéant les arrangements pris pour la répartition des recettes et des dépenses": financing,
        "Totaux_2": _fmt_amount(grand),
        "le_12": _fmt_date(timezone.localdate()),
        "Certifié exact_9": _officer_name(club, [CommitteeMandate.Role.PRESIDENT]),
    }
    _put_travel(
        values,
        mode,
        case.travel_units,
        case.travel_rate,
        travel_total,
        {
            "train": ("par train_2", "x_9", "1_29"),
            "car": ("par voiture_2", "km à_2", "2_30"),
            "plane": ("par avion_2", "x_10", "3_23"),
        },
    )
    if stay_total:
        values["Frais de séjour"] = str(case.stay_people or "")
        values["x_11"] = str(case.stay_days or "")
        values["jours à_2"] = _fmt_amount(_amount(case.stay_rate))
        values["1_31"] = _fmt_amount(stay_total)
    desc_fields = ["4_13", "5_8"]
    amount_fields = ["3_25", "4_14"]
    for index, (label, amount) in enumerate(divers[:2]):
        values[desc_fields[index]] = label
        values[amount_fields[index]] = _fmt_amount(amount)
    if len(divers) > 2:
        extra = "; ".join(f"{label} {_fmt_amount(amount)}" for label, amount in divers[2:])
        key = "Indiquez le mode de financement des rencontres et le cas échéant les arrangements pris pour la répartition des recettes et des dépenses"
        values[key] = "\n".join(part for part in (values.get(key), extra) if part)
    return values


def _organisation_values(case, athletes, officials, athlete_names, notes, travel_total, stay_total, divers, grand, president):
    remarks = notes
    if athlete_names:
        remarks = "\n".join(part for part in (remarks, "Athletes: " + ", ".join(athlete_names)) if part)
    labeled = []
    values = {
        "A RENSEIGNEMENTS GENERAUX_2": case.title,
        "Motivation de l’organisation": case.notes,
        "1_34": case.place,
        "2_35": " – ".join(part for part in (_fmt_date(case.starts_on), _fmt_date(case.ends_on)) if part),
        "1_35": "1" if athletes else "",
        "athlètes_2": str(len(athletes)) if athletes else "",
        "officiels": str(len(officials)) if officials else "",
        "le_14": _fmt_date(timezone.localdate()),
        "Le Président ou son représentant_7": president,
    }
    if travel_total:
        values["Dépenses 1"] = _fmt_amount(travel_total)
        values["Résultat 1"] = _fmt_amount(travel_total)
        labeled.append(f"Transport {_fmt_amount(travel_total)}")
    if stay_total:
        values["Dépenses 2"] = _fmt_amount(stay_total)
        values["Résultat 2"] = _fmt_amount(stay_total)
        labeled.append(f"Logement {_fmt_amount(stay_total)}")
    by_label = {label: amount for label, amount in divers}
    if by_label.get("Fournitures"):
        values["Dépenses 4"] = _fmt_amount(by_label["Fournitures"])
        values["Résultat 4"] = _fmt_amount(by_label["Fournitures"])
        labeled.append(f"Materiel {_fmt_amount(by_label['Fournitures'])}")
    if by_label.get("Inscription"):
        values["Dépenses 5"] = _fmt_amount(by_label["Inscription"])
        values["Résultat 5"] = _fmt_amount(by_label["Inscription"])
        labeled.append(f"Inscription {_fmt_amount(by_label['Inscription'])}")
    if by_label.get("Frais medicaux"):
        values["Dépenses 6"] = _fmt_amount(by_label["Frais medicaux"])
        values["Résultat 6"] = _fmt_amount(by_label["Frais medicaux"])
        labeled.append(f"Frais medicaux {_fmt_amount(by_label['Frais medicaux'])}")
    if grand:
        values["xxx_2"] = _fmt_amount(grand)
        values["Résultat 13"] = _fmt_amount(grand)
    if labeled:
        remarks = "\n".join(part for part in (remarks, "Devis: " + "; ".join(labeled)) if part)
    values["Remarques"] = remarks
    return values
