from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from pypdf import PdfReader
from rest_framework.test import APIClient

from accounts.models import User
from clubs.models import Club
from licenses.models import Income, License, LicenseType, Order
from members.models import Member
from modules.codes import build_payload, sign_payload
from modules.entitlements import install_id_str, redeem_product_code
from modules.registry import CLUB_MANAGEMENT_MODULE_ID

from .models import CoachQualification, Committee, CommitteeMandate


@override_settings(DEBUG=True, MODULE_CODE_PUBLIC_KEY="", MODULE_CODE_PRIVATE_KEY="")
class SubsidyTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_user(username="subsidy-admin", password="pass12345", role=User.Roles.CLUB_ADMIN)
        self.coach = User.objects.create_user(
            username="subsidy-coach", password="pass12345", role=User.Roles.COACH, first_name="Kim", last_name="Coach"
        )
        self.club = Club.objects.create(name="Subsidy Club", created_by=self.admin, iban="LU280019400644750000")
        self.club.admins.add(self.admin)
        self.club.trainers.add(self.coach)
        self.youth = Member.objects.create(
            club=self.club,
            first_name="Ada",
            last_name="Youth",
            date_of_birth=date(2012, 5, 1),
            sex=Member.Sex.FEMALE,
        )
        license_type = LicenseType.objects.create(name="Subsidy athlete", code="ath-sub")
        License.objects.create(
            member=self.youth,
            club=self.club,
            license_type=license_type,
            year=2026,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            status=License.Status.ACTIVE,
        )
        token = sign_payload(build_payload(module_ids=[CLUB_MANAGEMENT_MODULE_ID], install_id=install_id_str()))
        redeem_product_code(token)
        from modules.models import ClubModuleAssignment

        ClubModuleAssignment.objects.create(club=self.club, module_id=CLUB_MANAGEMENT_MODULE_ID, enabled=True)
        self.client.force_authenticate(user=self.admin)

    def test_dossier_counts_youth_and_coach_points(self):
        saved = self.client.post(
            f"/api/club-management/subsidies/coaches/?club={self.club.id}&year=2026",
            {"user_id": self.coach.id, "eqf_level": "eqf3", "coaches_under_16": True, "diploma_status": "attached"},
            format="json",
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual(saved.data["eligible_youth"], 1)
        self.assertEqual(saved.data["qualite_estimate"], "150.00")
        self.assertEqual(saved.data["coach_points"], 40)
        self.assertTrue(CoachQualification.objects.filter(user=self.coach, eqf_level="eqf3").exists())
        checks = {row["id"]: row["ok"] for row in saved.data["checklist"]}
        self.assertTrue(checks["youth"])
        self.assertTrue(checks["eqf3"])
        assistant = User.objects.create_user(
            username="subsidy-assistant",
            password="pass12345",
            role=User.Roles.COACH,
            first_name="Ann",
            last_name="Aide",
        )
        self.club.trainers.add(assistant)
        CoachQualification.objects.create(
            club=self.club,
            user=assistant,
            eqf_level=CoachQualification.Level.EQF1,
            include_in_qualite=False,
        )
        dossier = self.client.get(f"/api/club-management/subsidies/?club={self.club.id}&year=2026")
        names = [row["name"] for row in dossier.data["coaches"]]
        self.assertIn("Kim Coach", names)
        self.assertNotIn("Ann Aide", names)
        self.assertEqual(dossier.data["coach_points"], 40)

    def test_extraordinary_form_pdf(self):
        created = self.client.post(
            f"/api/club-management/subsidies/cases/?club={self.club.id}&year=2026",
            {
                "kind": "championship",
                "title": "European Championships",
                "place": "Sofia",
                "starts_on": "2026-05-01",
                "ends_on": "2026-05-04",
                "athlete_ids": [self.youth.id],
                "travel_units": "2",
                "travel_rate": "180",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 200, created.data)
        pdf = self.client.get(
            f"/api/club-management/subsidies/cases/{created.data['id']}/form.pdf?club={self.club.id}"
        )
        self.assertEqual(pdf.status_code, 200, getattr(pdf, "data", pdf.content[:120]))
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_paid_amount_waits_for_an_officer_then_joins_club_income(self):
        denied = self.client.patch(
            f"/api/club-management/subsidies/?club={self.club.id}&year=2026",
            {"status": "paid", "paid_amount": "400.00"},
            format="json",
        )
        self.assertEqual(denied.status_code, 200, denied.data)
        self.assertEqual(denied.data["season"]["submitted_on"], "")
        self.assertTrue(denied.data["season"]["paid_on"])
        self.assertTrue(denied.data["season"]["income_needs_officer"])
        self.assertFalse(Income.objects.filter(ledger=Order.Ledger.CLUB, club=self.club).exists())

        officer = Member.objects.create(
            club=self.club,
            user=self.admin,
            first_name="Rita",
            last_name="Treasurer",
            sex=Member.Sex.FEMALE,
        )
        committee = Committee.objects.create(club=self.club, scope=Committee.Scope.CLUB, name="Board")
        CommitteeMandate.objects.create(
            committee=committee,
            role=CommitteeMandate.Role.TREASURER,
            member=officer,
            started_on=date(2026, 1, 1),
        )
        opened = self.client.get(f"/api/club-management/subsidies/?club={self.club.id}&year=2026")
        self.assertEqual(opened.status_code, 200, opened.data)
        self.assertFalse(opened.data["season"]["income_needs_officer"])
        self.assertTrue(opened.data["season"]["income_number"])
        income = Income.objects.get(reference=f"subsidy-{self.club.id}-2026")
        self.assertEqual(income.amount, Decimal("400.00"))
        self.assertEqual(income.ledger, Order.Ledger.CLUB)
        self.assertEqual(income.payer, "Ministère des Sports")

    def test_rib_and_diploma_can_be_kept_for_myguichet(self):
        rib = SimpleUploadedFile("rib.pdf", b"%PDF-1.4 rib", content_type="application/pdf")
        saved = self.client.post(
            f"/api/club-management/subsidies/rib/?club={self.club.id}&year=2026",
            {"file": rib},
            format="multipart",
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertTrue(saved.data["season"]["has_rib"])
        checks = {row["id"]: row["ok"] for row in saved.data["checklist"]}
        self.assertTrue(checks["rib"])
        downloaded = self.client.get(f"/api/club-management/subsidies/rib/?club={self.club.id}&year=2026")
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn(b"%PDF", b"".join(downloaded.streaming_content))

        diploma = SimpleUploadedFile("diploma.pdf", b"%PDF-1.4 diploma", content_type="application/pdf")
        attached = self.client.post(
            f"/api/club-management/subsidies/coaches/{self.coach.id}/diploma/?club={self.club.id}&year=2026",
            {"file": diploma},
            format="multipart",
        )
        self.assertEqual(attached.status_code, 200, attached.data)
        coach = next(row for row in attached.data["coaches"] if row["user_id"] == self.coach.id)
        self.assertEqual(coach["diploma_status"], "attached")
        self.assertTrue(coach["has_diploma"])
        self.assertEqual(attached.data["dates"]["inaps_by"], "2026-12-15")
        self.assertEqual(attached.data["dates"]["diplomas_by"], "2027-02-01")

    def test_extraordinary_request_is_year_scoped_and_prints_costs(self):
        created = self.client.post(
            f"/api/club-management/subsidies/cases/?club={self.club.id}&year=2026",
            {
                "kind": "championship",
                "title": "European Championships",
                "place": "Sofia",
                "starts_on": "2026-05-01",
                "ends_on": "2026-05-04",
                "athlete_ids": [self.youth.id],
                "travel_mode": "train",
                "travel_units": "2",
                "travel_rate": "180",
                "entry_fee": "25",
            },
            format="json",
        )
        self.assertEqual(created.status_code, 200, created.data)
        other_year = self.client.get(f"/api/club-management/subsidies/?club={self.club.id}&year=2025")
        self.assertEqual(other_year.data["cases"], [])
        self.assertEqual(other_year.data["dates"]["diplomas_by"], "2026-02-01")
        case_id = created.data["id"]
        updated = self.client.patch(
            f"/api/club-management/subsidies/cases/{case_id}/?club={self.club.id}",
            {"status": "accounted", "place": "Sofia"},
            format="json",
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["cases"][0]["status"], "accounted")
        pdf = self.client.get(f"/api/club-management/subsidies/cases/{case_id}/form.pdf?club={self.club.id}")
        reader = PdfReader(BytesIO(pdf.content))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        folded = text.casefold()
        self.assertIn("youth ada", folded)
        self.assertIn("Sofia", text)
        self.assertIn("360.00", text)
        self.assertIn("Inscription", text)
        self.assertIn("25.00", text)
        removed = self.client.delete(f"/api/club-management/subsidies/cases/{case_id}/?club={self.club.id}")
        self.assertEqual(removed.status_code, 200, removed.data)
        self.assertEqual(removed.data["cases"], [])

    def test_effectifs_xlsx_follows_the_myguichet_page(self):
        senior = Member.objects.create(
            club=self.club,
            first_name="Omar",
            last_name="Senior",
            date_of_birth=date(2000, 1, 1),
            sex=Member.Sex.MALE,
            primary_license_role=Member.LicenseRole.ATHLETE,
        )
        coach = Member.objects.create(
            club=self.club,
            first_name="Lea",
            last_name="Coach",
            date_of_birth=date(1980, 3, 1),
            sex=Member.Sex.FEMALE,
            primary_license_role=Member.LicenseRole.COACH,
        )
        license_type = LicenseType.objects.get(code="ath-sub")
        left = Member.objects.create(
            club=self.club,
            first_name="Max",
            last_name="Left",
            date_of_birth=date(2014, 6, 1),
            sex=Member.Sex.MALE,
            primary_license_role=Member.LicenseRole.ATHLETE,
            is_active=False,
        )
        for member in (senior, coach, left):
            License.objects.create(
                member=member,
                club=self.club,
                license_type=license_type,
                year=2026,
                start_date=date(2026, 1, 1),
                end_date=date(2026, 12, 31),
                status=License.Status.ACTIVE,
            )
        self.client.patch(
            f"/api/club-management/subsidies/?club={self.club.id}&year=2026",
            {"non_licensed_count": 2},
            format="json",
        )
        response = self.client.get(f"/api/club-management/subsidies/youth.xlsx?club={self.club.id}&year=2026")
        self.assertEqual(response.status_code, 200, response.content[:120])
        self.assertTrue(response.content.startswith(b"PK"))
        workbook = load_workbook(BytesIO(response.content))
        effectifs = workbook["Effectifs"]
        labels = {effectifs.cell(row, 1).value: row for row in range(6, 24)}
        under = labels["Jeunes < 16 ans"]
        self.assertEqual(effectifs.cell(under, 2).value, 1)
        self.assertEqual(effectifs.cell(under, 3).value, 1)
        seniors = labels["Seniors (18–34)"]
        self.assertEqual(effectifs.cell(seniors, 2).value, 1)
        masters = labels["Vétérans / Masters (35+)"]
        self.assertIsNone(effectifs.cell(masters, 2).value)
        self.assertEqual(effectifs.cell(masters, 3).value, 1)
        coaches = labels["Entraîneurs / Moniteurs"]
        self.assertEqual(effectifs.cell(coaches, 3).value, 1)
        self.assertEqual(effectifs.cell(labels["Effectifs non licenciés du club"], 4).value, 2)
        names = workbook["Jeunes moins de 16"]
        youth_names = {
            names.cell(row, 1).value: names.cell(row, 2).value
            for row in range(2, names.max_row + 1)
        }
        self.assertEqual(youth_names["YOUTH"], "Ada")
        self.assertEqual(youth_names["LEFT"], "Max")

    def test_countersigned_trainers_list(self):
        Member.objects.create(
            club=self.club,
            user=self.coach,
            first_name="Faustino",
            last_name="Cima",
            belt_rank="1er dan",
            sex=Member.Sex.MALE,
        )
        CoachQualification.objects.create(
            club=self.club,
            user=self.coach,
            eqf_level=CoachQualification.Level.EQF3,
            coaches_under_16=True,
        )
        president_user = User.objects.create_user(
            username="subsidy-president",
            password="pass12345",
            role=User.Roles.CLUB_ADMIN,
            first_name="Patrick",
            last_name="Muller",
        )
        president = Member.objects.create(
            club=self.club,
            user=president_user,
            first_name="Patrick",
            last_name="Muller",
            sex=Member.Sex.MALE,
        )
        committee = Committee.objects.create(club=self.club, scope=Committee.Scope.CLUB, name="Board")
        CommitteeMandate.objects.create(
            committee=committee,
            role=CommitteeMandate.Role.PRESIDENT,
            member=president,
            started_on=date(2026, 1, 1),
        )
        self.club.locality = "Vichten"
        self.club.save(update_fields=["locality"])
        response = self.client.get(
            f"/api/club-management/subsidies/trainers.pdf?club={self.club.id}&year=2026"
        )
        self.assertEqual(response.status_code, 200, response.content[:180])
        self.assertTrue(response.content.startswith(b"%PDF"))
        text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn("CIMA Faustino", text)
        self.assertIn("1er dan", text)
        self.assertIn("EQF 3", text)
        self.assertIn("MULLER Patrick", text)
        self.assertIn("Vichten", text)
        self.assertIn("Liste d'entraîneurs contresignée", text)
        self.assertIn("Année de subside 2026", text)
        self.assertIn("Certifié exact", text)
        self.assertIn("certifie l'exactitude", text)

    def test_qualite_training_list_keeps_under_16_classes(self):
        from .training_models import TrainingSeries, TrainingSettings

        TrainingSettings.objects.create(club=self.club, default_place="Hall de Vichten")
        CoachQualification.objects.create(
            club=self.club,
            user=self.coach,
            eqf_level=CoachQualification.Level.EQF3,
            coaches_under_16=True,
        )
        kids = TrainingSeries.objects.create(
            club=self.club,
            name="Enfants mardi",
            audience=TrainingSeries.Audience.KIDS,
            weekday=1,
            start_time="17:00",
            end_time="18:30",
            valid_from=date(2026, 9, 1),
            valid_until=date(2027, 6, 30),
            counts_for_under_16=True,
            skip_public_holidays=True,
            skip_school_holidays=True,
        )
        kids.coaches.add(self.coach)
        kids.regulars.add(self.youth)
        adults = TrainingSeries.objects.create(
            club=self.club,
            name="Adultes jeudi",
            audience=TrainingSeries.Audience.ADULTS,
            weekday=3,
            start_time="19:00",
            end_time="20:30",
            valid_from=date(2026, 9, 1),
            valid_until=date(2027, 6, 30),
            counts_for_under_16=False,
        )
        response = self.client.get(f"/api/club-management/subsidies/trainings.pdf?club={self.club.id}&year=2026")
        self.assertEqual(response.status_code, 200, response.content[:180])
        text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn("Liste des entraînements", text)
        self.assertIn("Année de subside 2026", text)
        self.assertIn("Plages horaires", text)
        self.assertIn("Mardi", text)
        self.assertIn("Jeudi", text)
        self.assertNotIn("Lundi", text)
        self.assertNotIn("Dimanche", text)
        self.assertIn("Enfants", text)
        self.assertIn("17:00", text)
        self.assertIn("COACH", text)
        self.assertIn("Adultes", text)
        self.assertIn("19:00", text)
        self.assertIn("Certifié exact", text)
        self.assertIn("certifie l'exactitude", text)
        self.assertNotIn("Hall de", text)
        saved = self.client.patch(
            f"/api/club-management/subsidies/?club={self.club.id}&year=2026",
            {"training_place": "Salle B"},
            format="json",
        )
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual(saved.data["training_place"], "Salle B")

    def test_seventh_trainer_starts_the_next_page_with_the_president(self):
        for index in range(6):
            account = User.objects.create_user(
                username=f"subsidy-extra-coach-{index}",
                password="pass12345",
                first_name=f"Coach{index}",
                last_name=f"Extra{index}",
            )
            Member.objects.create(
                club=self.club,
                user=account,
                first_name=f"Coach{index}",
                last_name=f"Extra{index}",
                sex=Member.Sex.MALE,
            )
            self.club.trainers.add(account)
        response = self.client.get(
            f"/api/club-management/subsidies/trainers.pdf?club={self.club.id}&year=2026"
        )
        self.assertEqual(response.status_code, 200, response.content[:180])
        text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn("EXTRA0", text)
        self.assertIn("EXTRA5", text)
        self.assertIn("président", text)
        self.assertIn("Certifié exact", text)
        self.assertIn("Liste d'entraîneurs contresignée", text)

    def test_subsidy_lists_put_higher_grades_first(self):
        from .training_models import TrainingSeries

        senior = User.objects.create_user(
            username="subsidy-grade-dan",
            password="pass12345",
            first_name="High",
            last_name="Dan",
        )
        junior = User.objects.create_user(
            username="subsidy-grade-kup",
            password="pass12345",
            first_name="Low",
            last_name="Kup",
        )
        Member.objects.create(
            club=self.club, user=senior, first_name="High", last_name="Dan", belt_rank="4th Dan", sex=Member.Sex.MALE
        )
        Member.objects.create(
            club=self.club, user=junior, first_name="Low", last_name="Kup", belt_rank="1er kup", sex=Member.Sex.MALE
        )
        self.club.trainers.add(senior, junior)
        series = TrainingSeries.objects.create(
            club=self.club,
            name="Competition",
            audience=TrainingSeries.Audience.COMPETITION,
            weekday=0,
            start_time="18:00",
            end_time="19:30",
            valid_from=date(2026, 1, 1),
            valid_until=date(2026, 12, 31),
        )
        series.coaches.add(junior, senior)
        trainers = self.client.get(f"/api/club-management/subsidies/trainers.pdf?club={self.club.id}&year=2026")
        training = self.client.get(f"/api/club-management/subsidies/trainings.pdf?club={self.club.id}&year=2026")
        trainers_text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(trainers.content)).pages)
        training_text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(training.content)).pages)
        self.assertLess(trainers_text.index("DAN High"), trainers_text.index("KUP Low"))
        self.assertLess(training_text.index("DAN High"), training_text.index("KUP Low"))

    def test_extraordinary_officials_are_adults_only(self):
        adult = Member.objects.create(
            club=self.club,
            first_name="Rita",
            last_name="Adult",
            date_of_birth=date(1990, 1, 1),
            sex=Member.Sex.FEMALE,
        )
        created = self.client.post(
            f"/api/club-management/subsidies/cases/?club={self.club.id}&year=2026",
            {
                "kind": "championship",
                "title": "Open",
                "official_ids": [self.youth.id, adult.id],
            },
            format="json",
        )
        self.assertEqual(created.status_code, 200, created.data)
        self.assertEqual(created.data["cases"][0]["official_ids"], [adult.id])
        officials = [row for row in created.data["members"] if row["adult"]]
        self.assertIn(adult.id, [row["id"] for row in officials])
        self.assertNotIn(self.youth.id, [row["id"] for row in officials])
